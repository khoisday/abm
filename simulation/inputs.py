"""Inspection of the source files and rule-registry extraction.

Only facts actually read from the files are written. Interpretations are marked
as interpretations; draft self-assessment answers are not treated as verified
facts or as proof of guest compliance.
"""
import csv
from pathlib import Path

import pandas as pd

from .config import ROOT

EXCEL = "assessment_answers_categorized.xlsx"
SOURCE_CONFIG = ROOT / "config" / "source_data.json"   # local only (not versioned): {"facility": "<name in the workbook>"}


def target_facility():
    """Name of the assessed property inside the workbook; kept out of version control."""
    import json
    import os
    if os.environ.get("SOURCE_FACILITY"):
        return os.environ["SOURCE_FACILITY"]
    if SOURCE_CONFIG.exists():
        return json.loads(SOURCE_CONFIG.read_text(encoding="utf-8")).get("facility")
    return None

# code -> (category, interpretation, agent/equipment, trigger, action, parameters, status, model_link)
RULES = {
    "Q1.1": ("static_building_info", "Tự khai 533 phòng – trùng với yêu cầu người dùng (533).", "RoomBank", "-", "-",
             "hotel.rooms", "implemented", "Số phòng mặc định = 533 (yêu cầu người dùng)."),
    "Q1.2": ("static_building_info", "Lượt phòng/năm tự khai; 132776/(533×365)=68,2% – không khớp Q1.3 (53%).",
             "-", "-", "-", "-", "reference_only", "Dùng làm tham chiếu; không hiệu chỉnh."),
    "Q1.3": ("static_building_info", "Công suất phòng tự khai 53%; mô hình mặc định dùng 80% theo đề cương dự án.",
             "-", "-", "-", "hotel.occupancy", "reference_only", "Người dùng có thể đặt 0,53 trong dashboard."),
    "Q1.4": ("static_building_info", "Diện tích sưởi/làm mát.", "-", "-", "-", "-", "not_applicable",
             "Mô hình phòng dùng UA/C tổng hợp, không theo m2."),
    "Q1.5": ("static_building_info", "Diện tích điều hòa phòng & hành lang.", "-", "-", "-", "-", "not_applicable", ""),
    "Q1.6": ("static_building_info", "Diện tích điều hòa khu chung.", "-", "-", "-", "-", "not_applicable", ""),
    "Q1.7": ("static_building_info", "Tổng diện tích sàn.", "-", "-", "-", "-", "not_applicable", ""),
    "Q1.8": ("assessment_answer", "Điện lưới theo tháng (kWh) – tự khai, bản nháp. Không rõ ranh giới (có gồm công viên?).",
             "-", "-", "-", "-", "reference_only", "Hiển thị như tham chiếu, không cộng vào điện bottom-up."),
    "Q1.20": ("assessment_answer", "Điện mặt trời theo tháng (kWh) – tự khai.", "-", "-", "-", "-", "deferred",
              "Chưa mô hình hóa PV tại chỗ."),
    "Q2.13": ("existing_practice", "Tỷ lệ LED 76–100%.", "RoomBank lighting", "-", "-", "rooms.loads_kw",
              "implemented", "Tải chiếu sáng phòng giả định mức thấp (LED)."),
    "Q2.14": ("existing_practice", "Cảm biến HVAC tự động phòng lưu trú 76–100%.", "RoomBank HVAC",
              "Khách rời phòng", "Ngắt/giảm điều hòa", "rooms.keycard_bypass_share", "implemented",
              "Baseline: thẻ từ ngắt HVAC; 25% phòng bị 'bypass' (giả định, nhất quán với 76–100%)."),
    "Q2.15": ("existing_practice", "Cảm biến HVAC tự động khu chung 76–100%.", "hotel_common", "-", "-", "-",
              "not_applicable", "Khu chung là tải tổng hợp theo lịch."),
    "Q2.16": ("existing_practice", "Rèm cách nhiệt.", "RoomBank", "-", "-", "rooms.solar_gain_peak_kw",
              "implemented", "Bức xạ mặt trời vào phòng giả định thấp."),
    "Q2.17": ("existing_practice", "Chế độ tiết kiệm năng lượng cho thiết bị.", "RoomBank plug", "-", "-", "-",
              "not_applicable", ""),
    "Q2.18": ("actionable_rule", "Nhiệt độ điều hòa phòng khách 24–26°C (mùa hè); nước nóng 50–70°C.",
              "RoomBank HVAC", "Phòng có khách", "Điều khiển theo điểm đặt", "rooms.guest_setpoints_c",
              "implemented", "Điểm đặt khách 23–26°C (đa số 24–25°C); chính sách ≠ tuân thủ của khách."),
    "Q2.42": ("existing_practice", "Dọn phòng hằng ngày; thay ga/khăn chỉ khi khách yêu cầu.", "Housekeeper",
              "Dọn phòng ở lại", "Thay khăn nếu khách yêu cầu", "housekeeping.towel_request_base_prob",
              "implemented", "Xác suất yêu cầu thay khăn là giả định; không coi là tuân thủ 100%."),
    "Q2.36": ("existing_practice", "Tái sử dụng ga/gối/khăn cũ.", "-", "-", "-", "-", "not_applicable", ""),
    "Q2.23": ("existing_practice", "Thiết bị giặt là tiết kiệm nước.", "Laundry", "-", "-", "laundry", "deferred",
              "Nước không được mô hình hóa."),
    "Q2.39": ("existing_practice", "Dùng phương tiện xanh chở khách.", "Shuttle", "-", "-",
              "transport.include_transport_energy", "deferred", "Điện phương tiện mặc định nằm ngoài ranh giới."),
    "Q2.40": ("existing_practice", "Trạm sạc xe điện.", "-", "-", "-", "-", "not_applicable", ""),
    "Q2.41": ("existing_practice", "Xe đạp/xe đạp điện cho khách.", "-", "-", "-", "-", "deferred", ""),
    "Q2.79": ("aspirational_policy", "Chính sách khuyến khích khách/NV tiết kiệm tài nguyên.", "GuestGroup",
              "Gợi ý", "Khách có thể từ chối", "rec_acceptance", "implemented",
              "Dùng như cơ sở cho kịch bản gợi ý & khuyến khích tái sử dụng khăn (không giả định tuân thủ)."),
    "Q2.7": ("existing_practice", "Theo dõi năng lượng hằng tuần.", "-", "-", "-", "-", "not_applicable", ""),
    "Q2.9": ("existing_practice", "Đồng hồ điện theo bộ phận.", "Meter", "-", "-", "-", "not_applicable",
             "Gợi ý dữ liệu theo bộ phận có thể dùng hiệu chỉnh trong tương lai."),
    "Q2.11": ("existing_practice", "Hệ thống giám sát năng lượng tập trung.", "-", "-", "-", "-", "not_applicable", ""),
    "Q2.5": ("existing_practice", "Đo phát thải carbon hằng tháng.", "-", "-", "-", "-", "not_applicable", ""),
    "Q2.6": ("aspirational_policy", "Có kế hoạch giảm phát thải.", "-", "-", "-", "-", "not_applicable", ""),
}


def build_rule_registry(xlsx=ROOT / EXCEL, out=ROOT / "rule_registry.csv"):
    import openpyxl
    facility = target_facility()
    if facility is None:
        raise SystemExit(f"Set the assessed property name in {SOURCE_CONFIG} or SOURCE_FACILITY to rebuild the registry")
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets[1:]:
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == 1 or row[0] != facility:
                continue
            code, indicator, qtype, unit, value = row[3], row[4], row[5], row[6], row[7]
            months = [v for v in row[8:20]]
            if value is None and any(m is not None for m in months):
                value = "; ".join(f"T{k + 1}={m}" for k, m in enumerate(months))
            cat, interp, agent, trig, act, params, status, link = RULES.get(
                code, ("assessment_answer", "Câu trả lời đánh giá, không phải quy tắc hành vi/thiết bị thực thi được.",
                       "-", "-", "-", "-", "not_applicable", ""))
            rows.append(dict(rule_id=f"HTL-{code}", source_file=EXCEL, sheet=ws.title, row=i, code=code,
                             original_text=f"{indicator}: {value}" + (f" [{unit}]" if unit else ""),
                             question_type=qtype, record_status=row[2], review_status=row[20],
                             category=cat, interpretation=interp, agent_or_equipment=agent, trigger=trig,
                             action=act, required_parameters=params, status=status, model_link=link))
    df = pd.DataFrame(rows)
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return df


def self_reported_monthly_grid_kwh(xlsx=ROOT / EXCEL):
    """Self-reported (draft, pending) monthly grid electricity; reference only. None if the file is absent."""
    import openpyxl
    facility = target_facility()
    if not xlsx.exists() or facility is None:
        return None
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row[0] == facility and row[3] == "Q1.8":
                return dict(sheet=ws.title, row=i, kwh=[float(v) for v in row[8:20]])
    return None


def find_reference_csv():
    hits = sorted(ROOT.glob("*_15min_NREL.csv")) + sorted((ROOT / "data" / "imported").glob("*_15min_NREL.csv"))
    return hits[0] if hits else None


def inspect_reference_csv(path=None):
    path = Path(path or find_reference_csv())
    df = pd.read_csv(path)
    ts = pd.to_datetime(df["timestamp"])
    elec_cols = [c for c in df.columns if c.startswith("out.electricity.") and c.endswith(".energy_consumption")]
    enduse = [c for c in elec_cols if ".total." not in c]
    info = dict(file=path.name, rows=len(df), bldg_ids=sorted(df["bldg_id"].astype(str).unique().tolist()),
                first_timestamp=str(ts.iloc[0]), last_timestamp=str(ts.iloc[-1]),
                interval_minutes=sorted(set((ts.diff().dropna().dt.total_seconds() / 60).astype(int)))[:5],
                duplicate_timestamps=int(ts.duplicated().sum()), missing_cells=int(df.isna().sum().sum()),
                electricity_enduse_columns=[c.split(".")[2] for c in enduse],
                annual_electricity_kwh=float(df["out.electricity.total.energy_consumption"].sum()),
                enduse_sum_matches_total=bool(abs(df[enduse].sum(axis=1).sum()
                                                  - df["out.electricity.total.energy_consumption"].sum()) < 5))
    return info


def reference_daily_profile(path=None, month=6):
    """Mean 15-minute electricity profile (kW) for a month, normalized to its daily mean.

    Timestamps are treated as interval-ending (first stamp 00:15, last 00:00 next year);
    each value is kWh per 15 min, so kW = kWh / 0.25. Used for shape comparison only.
    """
    path = Path(path or find_reference_csv())
    df = pd.read_csv(path, usecols=["timestamp", "out.electricity.total.energy_consumption"])
    ts = pd.to_datetime(df["timestamp"]) - pd.Timedelta(minutes=15)   # interval start
    df["kw"] = df["out.electricity.total.energy_consumption"] / 0.25
    df = df[ts.dt.month == month]
    slot = (ts[df.index].dt.hour * 4 + ts[df.index].dt.minute // 15)
    prof = df.groupby(slot.values)["kw"].mean()
    return (prof / prof.mean()).reset_index(drop=True)


if __name__ == "__main__":
    reg = build_rule_registry()
    print(f"rule_registry.csv: {len(reg)} rows; implemented={int((reg.status == 'implemented').sum())}")
    print(inspect_reference_csv())
